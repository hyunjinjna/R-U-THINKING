# -*- coding: utf-8 -*-
# 파닉스 게이트 5단계 × 4문항 = 20문항 — 전부 듣기형 (Julia 확정: 한국어 보기 제거)
# 음성파일 열 = TTS로 읽을 단어 (재생 버튼은 다음 코딩에서 추가 — 그 전엔 투입 불가)
# 1단계: 단어 듣고 첫 글자 고르기 / 2~5단계: 단어 듣고 맞는 철자 고르기
# 단계: 1 글자 소리 / 2 단모음 CVC / 3 장모음(magic e·이중모음) / 4 자음 블렌드 / 5 이중글자(sh·ch·th·ph)
import openpyxl
from openpyxl.styles import Font

Q_LETTER = '소리를 듣고 첫 글자를 고르세요.'
Q_WORD = '소리를 듣고 맞는 단어를 고르세요.'

# (단계, 질문, 정답, [오답3], TTS단어)
Q = [
# 1단계: 첫 글자 소리 (b/d/p 모양 혼동, 소리 유사 글자 함정)
(1, Q_LETTER, 'b', ['d','p','v'], 'ball'),
(1, Q_LETTER, 'm', ['n','w','h'], 'moon'),
(1, Q_LETTER, 's', ['z','c','j'], 'sun'),
(1, Q_LETTER, 'g', ['j','k','d'], 'girl'),
# 2단계: 단모음 CVC (모음 소리 구분이 핵심)
(2, Q_WORD, 'cat', ['cut','kit','cot'], 'cat'),
(2, Q_WORD, 'pen', ['pin','pan','pun'], 'pen'),
(2, Q_WORD, 'hot', ['hat','hut','hit'], 'hot'),
(2, Q_WORD, 'bed', ['bad','bud','bid'], 'bed'),
# 3단계: 장모음·magic e (cake/cap, bike/big이 magic e 판별)
(3, Q_WORD, 'rain', ['ran','run','rail'], 'rain'),
(3, Q_WORD, 'cute', ['cut','coat','kit'], 'cute'),
(3, Q_WORD, 'cake', ['cap','kick','coke'], 'cake'),
(3, Q_WORD, 'bike', ['big','back','beak'], 'bike'),
# 4단계: 자음 블렌드 (블렌드 유무·자리 구분)
(4, Q_WORD, 'stop', ['top','shop','sop'], 'stop'),
(4, Q_WORD, 'frog', ['fog','log','flag'], 'frog'),
(4, Q_WORD, 'hand', ['had','ham','and'], 'hand'),
(4, Q_WORD, 'milk', ['mill','melt','mix'], 'milk'),
# 5단계: 이중글자 sh·ch·th·ph
(5, Q_WORD, 'ship', ['sip','chip','sheep'], 'ship'),
(5, Q_WORD, 'chin', ['shin','tin','thin'], 'chin'),
(5, Q_WORD, 'think', ['sink','thin','pink'], 'think'),
(5, Q_WORD, 'phone', ['bone','foam','fan'], 'phone'),
]

# 검증
levels = {}
for lv, q, ans, ds, tts in Q:
    levels[lv] = levels.get(lv, 0) + 1
    assert len(set([ans] + ds)) == 4, q + ans
    assert tts, ans
    if q == Q_WORD:
        assert ans == tts, ans           # 듣는 단어 = 정답 철자
    else:
        assert tts.startswith(ans), ans  # 첫 글자 문항: TTS 단어의 첫 글자 = 정답
assert len(levels) == 5 and all(v == 4 for v in levels.values()), levels

wb = openpyxl.Workbook()
ws = wb.active
ws.title = '파닉스'
cols = ['문제ID','게이트단계(1-5)','질문유형','질문','보기1','보기2','보기3','보기4','정답번호','이미지파일','음성파일']
ws.append(cols)
for c in ws[1]: c.font = Font(name='Arial', bold=True)

counter = {}
pos = 0
for lv, q, ans, ds, tts in Q:
    counter[lv] = counter.get(lv, 0) + 1
    qid = f"PHON-{lv}-{counter[lv]:03d}"
    answer_idx = pos % 4
    pos += 1
    opts = [None]*4
    opts[answer_idx] = ans
    di = 0
    for i in range(4):
        if opts[i] is None:
            opts[i] = ds[di]; di += 1
    ws.append([qid, lv, '듣고고르기', q] + opts + [answer_idx+1, '', tts])
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = Font(name='Arial')
ws.column_dimensions['D'].width = 40
ws.column_dimensions['A'].width = 14
ws.column_dimensions['C'].width = 12

wb.save('/mnt/user-data/outputs/파닉스_문제은행_20문항.xlsx')

from collections import Counter
apos = Counter(r[8] for r in ws.iter_rows(min_row=2, values_only=True))
print('문항 수:', ws.max_row - 1)
print('정답 위치:', dict(sorted(apos.items())))
print('전부 듣기형, 음성 텍스트 채움:', all(r[10] for r in ws.iter_rows(min_row=2, values_only=True)))
