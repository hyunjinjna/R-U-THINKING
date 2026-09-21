# -*- coding: utf-8 -*-
# 단어 영역 18단계 × 4문항 = 72문항
# 형식: (단계, 문장, 정답, [(오답, 구멍태그) × 3])
# 원칙: 문장에 힌트 내장 / 오답은 그 레벨 아이가 아는 실존 단어 / 정답 위치는 순환 분산
import openpyxl
from openpyxl.styles import Font

Q = [
# ===== Pre-A1-하 =====
("Pre-A1-하", 'Look at the sky. The _____ is yellow and hot.', "sun",
 [("son","유사단어혼동"),("moon","문맥추론부족"),("desk","단어뜻모름")]),
("Pre-A1-하", 'I see a big _____. It says "woof woof."', "dog",
 [("duck","유사단어혼동"),("cat","문맥추론부족"),("cup","단어뜻모름")]),
("Pre-A1-하", 'This is my _____. I write with it.', "pencil",
 [("people","유사단어혼동"),("book","문맥추론부족"),("milk","단어뜻모름")]),
("Pre-A1-하", 'I have two _____. I can see with them.', "eyes",
 [("ice","유사단어혼동"),("ears","문맥추론부족"),("bag","단어뜻모름")]),
# ===== Pre-A1-중 =====
("Pre-A1-중", 'It is raining. I need my _____.', "umbrella",
 [("uncle","유사단어혼동"),("sunglasses","문맥추론부족"),("table","단어뜻모름")]),
("Pre-A1-중", "My _____ is my father's mother.", "grandmother",
 [("grandfather","유사단어혼동"),("sister","문맥추론부족"),("window","단어뜻모름")]),
("Pre-A1-중", 'I am hungry. I want to _____ some bread.', "eat",
 [("ear","유사단어혼동"),("drink","문맥추론부족"),("jump","단어뜻모름")]),
("Pre-A1-중", 'The apple is red and the banana is _____.', "yellow",
 [("hello","유사단어혼동"),("blue","문맥추론부족"),("chair","단어뜻모름")]),
# ===== Pre-A1-상 =====
("Pre-A1-상", 'We play soccer in the _____ after school.', "playground",
 [("playtime","유사단어혼동"),("classroom","문맥추론부족"),("spoon","단어뜻모름")]),
("Pre-A1-상", 'Please open the window. The room is too _____.', "hot",
 [("hat","유사단어혼동"),("cold","문맥추론부족"),("green","단어뜻모름")]),
("Pre-A1-상", 'My little brother is only three years _____.', "old",
 [("gold","유사단어혼동"),("tall","문맥추론부족"),("fast","단어뜻모름")]),
("Pre-A1-상", 'I brush my _____ after eating candy.', "teeth",
 [("teach","유사단어혼동"),("hair","문맥추론부족"),("door","단어뜻모름")]),
# ===== A1-하 =====
("A1-하", 'Sam is _____ because his team lost the game.', "sad",
 [("said","유사단어혼동"),("happy","문맥추론부족"),("round","단어뜻모름")]),
("A1-하", 'Birds _____ high in the sky.', "fly",
 [("fry","유사단어혼동"),("swim","문맥추론부족"),("cook","단어뜻모름")]),
("A1-하", 'Put on your _____. Your feet are cold.', "socks",
 [("clocks","유사단어혼동"),("gloves","문맥추론부족"),("juice","단어뜻모름")]),
("A1-하", 'We buy milk and eggs at the _____.', "store",
 [("story","유사단어혼동"),("park","문맥추론부족"),("spoon","단어뜻모름")]),
# ===== A1-중 =====
("A1-중", 'The baby is sleeping. Please be _____.', "quiet",
 [("quite","유사단어혼동"),("loud","문맥추론부족"),("heavy","단어뜻모름")]),
("A1-중", "I can't find my keys. Will you help me _____ for them?", "look",
 [("book","유사단어혼동"),("listen","문맥추론부족"),("dance","단어뜻모름")]),
("A1-중", 'In winter, we make a _____ with snow.', "snowman",
 [("showman","유사단어혼동"),("sandcastle","문맥추론부족"),("lunch","단어뜻모름")]),
("A1-중", 'My grandma tells us funny _____ before bed.', "stories",
 [("stores","유사단어혼동"),("songs","문맥추론부족"),("shoes","단어뜻모름")]),
# ===== A1-상 =====
("A1-상", 'Turn off the light when you _____ the room.', "leave",
 [("live","유사단어혼동"),("enter","문맥추론부족"),("draw","단어뜻모름")]),
("A1-상", 'This box is too _____ for me. Can you carry it?', "heavy",
 [("happy","유사단어혼동"),("light","문맥추론부족"),("early","단어뜻모름")]),
("A1-상", "The train arrives at seven. Don't be _____!", "late",
 [("lake","유사단어혼동"),("early","문맥추론부족"),("soft","단어뜻모름")]),
("A1-상", 'She drew a _____ of her family in art class.', "picture",
 [("pitcher","유사단어혼동"),("photo","문맥추론부족"),("kitchen","단어뜻모름")]),
# ===== A2-하 =====
("A2-하", "The museum is _____ on Mondays, so we can't go in.", "closed",
 [("clothes","유사단어혼동"),("open","문맥추론부족"),("angry","단어뜻모름")]),
("A2-하", 'Dinosaurs lived a long time _____.', "ago",
 [("age","유사단어혼동"),("later","문맥추론부족"),("above","단어뜻모름")]),
("A2-하", 'The movie will _____ at 9 and end at 11.', "begin",
 [("behind","유사단어혼동"),("finish","문맥추론부족"),("borrow","단어뜻모름")]),
("A2-하", 'Wear a helmet to keep your head _____.', "safe",
 [("save","유사단어혼동"),("clean","문맥추론부족"),("round","단어뜻모름")]),
# ===== A2-중 =====
("A2-중", 'Everyone was excited about the _____ to the science museum.', "trip",
 [("trap","유사단어혼동"),("ticket","문맥추론부족"),("river","단어뜻모름")]),
("A2-중", 'Water _____ at 100 degrees.', "boils",
 [("builds","유사단어혼동"),("freezes","문맥추론부족"),("smiles","단어뜻모름")]),
("A2-중", 'He said sorry, and I decided to _____ him.', "forgive",
 [("forget","유사단어혼동"),("punish","문맥추론부족"),("measure","단어뜻모름")]),
("A2-중", 'The bridge _____ the two sides of the river.', "connects",
 [("collects","유사단어혼동"),("crosses","문맥추론부족"),("invites","단어뜻모름")]),
# ===== A2-상 =====
("A2-상", 'Please speak louder. Your voice is too _____ to hear.', "low",
 [("law","유사단어혼동"),("high","문맥추론부족"),("wide","단어뜻모름")]),
("A2-상", 'The scientist did an _____ to test her idea.', "experiment",
 [("experience","유사단어혼동"),("exercise","문맥추론부족"),("envelope","단어뜻모름")]),
("A2-상", 'We need to _____ energy by turning off lights.', "save",
 [("safe","유사단어혼동"),("waste","문맥추론부족"),("marry","단어뜻모름")]),
("A2-상", 'The road was _____ because of the heavy snow.', "blocked",
 [("black","유사단어혼동"),("smooth","문맥추론부족"),("polite","단어뜻모름")]),
# ===== B1-하 =====
("B1-하", 'Thank you for your _____. You really helped me.', "kindness",
 [("kind","품사형태변형부족"),("kingdom","유사단어혼동"),("lettuce","단어뜻모름")]),
("B1-하", 'The weather _____ says it will rain tomorrow.', "forecast",
 [("forehead","유사단어혼동"),("calendar","문맥추론부족"),("pillow","단어뜻모름")]),
("B1-하", 'You must _____ the rules of the game.', "follow",
 [("fellow","유사단어혼동"),("break","문맥추론부족"),("bake","단어뜻모름")]),
("B1-하", 'The two countries made peace after a long _____.', "war",
 [("wore","유사단어혼동"),("party","문맥추론부족"),("cousin","단어뜻모름")]),
# ===== B1-중 =====
("B1-중", 'The company will _____ a new phone next month.', "release",
 [("realize","유사단어혼동"),("repair","문맥추론부족"),("swallow","단어뜻모름")]),
("B1-중", 'His _____ to win made him practice every day.', "determination",
 [("determine","품사형태변형부족"),("destination","유사단어혼동"),("temperature","단어뜻모름")]),
("B1-중", 'The loud noise _____ the sleeping baby.', "disturbed",
 [("distributed","유사단어혼동"),("comforted","문맥추론부족"),("celebrated","단어뜻모름")]),
("B1-중", 'Please _____ your seatbelt during the flight.', "fasten",
 [("fashion","유사단어혼동"),("loosen","문맥추론부족"),("whisper","단어뜻모름")]),
# ===== B1-상 =====
("B1-상", 'The detective found an important _____ at the scene.', "clue",
 [("glue","유사단어혼동"),("crime","문맥추론부족"),("recipe","단어뜻모름")]),
("B1-상", 'Her _____ of the poem was different from mine.', "interpretation",
 [("interruption","유사단어혼동"),("interpret","품사형태변형부족"),("refrigerator","단어뜻모름")]),
("B1-상", 'The city plans to _____ the old library into a museum.', "transform",
 [("transport","유사단어혼동"),("destroy","문맥추론부족"),("apologize","단어뜻모름")]),
("B1-상", 'You should _____ your answers before handing in the test.', "review",
 [("preview","유사단어혼동"),("erase","문맥추론부족"),("boil","단어뜻모름")]),
# ===== B2-하 =====
("B2-하", 'The graph _____ a sharp increase in sales.', "indicates",
 [("dedicates","유사단어혼동"),("hides","문맥추론부족"),("swims","단어뜻모름")]),
("B2-하", 'Students must _____ their essays by Friday.', "submit",
 [("summit","유사단어혼동"),("receive","문맥추론부족"),("freeze","단어뜻모름")]),
("B2-하", 'The new law will _____ everyone in the city.', "affect",
 [("effect","유사단어혼동"),("defend","문맥추론부족"),("whistle","단어뜻모름")]),
("B2-하", 'Her argument was based on solid _____.', "evidence",
 [("evident","품사형태변형부족"),("avenue","유사단어혼동"),("laughter","단어뜻모름")]),
# ===== B2-중 =====
("B2-중", 'The two results were almost _____, with no real difference.', "identical",
 [("identity","품사형태변형부족"),("idiom","유사단어혼동"),("portable","단어뜻모름")]),
("B2-중", 'To raise money, the government decided to _____ new taxes.', "impose",
 [("expose","유사단어혼동"),("cancel","문맥추론부족"),("inhale","단어뜻모름")]),
("B2-중", 'He spoke with great _____ about climate change.', "passion",
 [("passive","유사단어혼동"),("silence","문맥추론부족"),("luggage","단어뜻모름")]),
("B2-중", 'The results may _____ depending on the weather.', "vary",
 [("very","유사단어혼동"),("stay","문맥추론부족"),("obey","단어뜻모름")]),
# ===== B2-상 =====
("B2-상", 'The committee will _____ the proposal before voting.', "evaluate",
 [("evacuate","유사단어혼동"),("ignore","문맥추론부족"),("dissolve","단어뜻모름")]),
("B2-상", 'His constant interruptions were a real _____.', "nuisance",
 [("nuance","유사단어혼동"),("pleasure","문맥추론부족"),("molecule","단어뜻모름")]),
("B2-상", 'The medicine can _____ the pain but not cure the disease.', "relieve",
 [("believe","유사단어혼동"),("increase","문맥추론부족"),("translate","단어뜻모름")]),
("B2-상", 'There was a striking _____ between the rich and poor areas.', "contrast",
 [("contract","유사단어혼동"),("similarity","문맥추론부족"),("ceiling","단어뜻모름")]),
# ===== C1-하 =====
("C1-하", 'Surprisingly, the findings _____ the original hypothesis.', "contradict",
 [("contract","유사단어혼동"),("support","문맥추론부족"),("memorize","단어뜻모름")]),
("C1-하", 'She has a _____ knowledge of European history.', "comprehensive",
 [("comprehension","품사형태변형부족"),("compressive","유사단어혼동"),("muddy","단어뜻모름")]),
("C1-하", 'The negotiations reached a _____ after both sides refused to move.', "deadlock",
 [("headlock","유사단어혼동"),("agreement","문맥추론부족"),("kettle","단어뜻모름")]),
("C1-하", "Shocking new evidence may _____ the court's original decision.", "overturn",
 [("overhear","유사단어혼동"),("confirm","문맥추론부족"),("celery","단어뜻모름")]),
# ===== C1-중 =====
("C1-중", 'His comments were _____ and offended many listeners.', "derogatory",
 [("directory","유사단어혼동"),("flattering","문맥추론부족"),("rectangular","단어뜻모름")]),
("C1-중", 'The species is on the _____ of extinction.', "verge",
 [("merge","유사단어혼동"),("center","문맥추론부족"),("syrup","단어뜻모름")]),
("C1-중", 'The report _____ the causes of the economic crisis in detail.', "analyzes",
 [("analysis","품사형태변형부족"),("paralyzes","유사단어혼동"),("whispers","단어뜻모름")]),
("C1-중", 'Efforts to _____ the conflict have failed so far.', "resolve",
 [("revolve","유사단어혼동"),("prolong","문맥추론부족"),("marinate","단어뜻모름")]),
# ===== C1-상 =====
("C1-상", "The politician's speech was full of _____ promises that meant nothing.", "hollow",
 [("hallow","유사단어혼동"),("sincere","문맥추론부족"),("gravitational","단어뜻모름")]),
("C1-상", 'The committee must _____ carefully before making such a costly decision.', "deliberate",
 [("deliver","유사단어혼동"),("rush","문맥추론부족"),("evaporate","단어뜻모름")]),
("C1-상", 'He examined the contract with _____ care.', "meticulous",
 [("meticulously","품사형태변형부족"),("momentous","유사단어혼동"),("edible","단어뜻모름")]),
("C1-상", 'The drought will _____ the food shortage in the region.', "exacerbate",
 [("exaggerate","유사단어혼동"),("relieve","문맥추론부족"),("hibernate","단어뜻모름")]),
]

# 검증: 18단계 × 4, 빈칸 존재, 보기 중복 없음
levels = {}
for lv, s, ans, ds in Q:
    levels[lv] = levels.get(lv, 0) + 1
    assert '_____' in s, s
    words = [ans] + [d[0] for d in ds]
    assert len(set(w.lower() for w in words)) == 4, s
assert len(levels) == 18 and all(v == 4 for v in levels.values()), levels

wb = openpyxl.Workbook()
ws = wb.active
ws.title = '단어'
cols = ['문제ID','CEFR단계','문장(빈칸포함)','보기1','보기2','보기3','보기4','정답번호','보기1구멍','보기2구멍','보기3구멍','보기4구멍']
ws.append(cols)
for c in ws[1]: c.font = Font(name='Arial', bold=True)

counter = {}
pos = 0  # 정답 위치 순환 (1→2→3→4)
for lv, sent, ans, ds in Q:
    counter[lv] = counter.get(lv, 0) + 1
    qid = f"VOCA-{lv.replace('-','')}-{counter[lv]:03d}"
    answer_idx = pos % 4
    pos += 1
    opts = [None] * 4
    tags = [''] * 4
    opts[answer_idx] = ans
    di = 0
    for i in range(4):
        if opts[i] is None:
            opts[i] = ds[di][0]
            tags[i] = ds[di][1]
            di += 1
    ws.append([qid, lv, sent] + opts + [answer_idx + 1] + tags)

for row in ws.iter_rows(min_row=2):
    for c in row: c.font = Font(name='Arial')
ws.column_dimensions['C'].width = 60
ws.column_dimensions['A'].width = 18
ws.column_dimensions['B'].width = 12

wb.save('/mnt/user-data/outputs/단어_문제은행_72문항.xlsx')

# 정답 위치 분포·태그 분포 확인
from collections import Counter
apos = Counter()
tcnt = Counter()
for r in ws.iter_rows(min_row=2, values_only=True):
    apos[r[7]] += 1
    for t in r[8:12]:
        if t: tcnt[t] += 1
print('문항 수:', ws.max_row - 1)
print('정답 위치 분포:', dict(apos))
print('구멍 태그 분포:', dict(tcnt))
