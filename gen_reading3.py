# -*- coding: utf-8 -*-
"""리딩 문제은행 보기 영어 전환 (2026-09-24 확정)
- 질문은 한국어 유지, 보기(선택지)만 전부 영어로 — 실전 원서 문제 형식
- 보기 영어 수준은 지문보다 한 단계 쉬움 (보기가 새 독해 문제가 되지 않게)
- 구멍 태그·정답 위치·지문·질문은 그대로 (셀 텍스트만 교체)
- 기존 산출물 /mnt/user-data/outputs/문제은행_리딩.xlsx 를 읽어 교체 후 재저장
"""
import openpyxl
import re

# 한국어 보기 → 영어 보기 (문항 설계 의도 보존 번역)
SWAP = {
    # READ-120-003 thirsty
    '목이 마른': 'wanting water',
    '배가 고픈': 'wanting food',
    '졸린': 'sleepy',
    '화가 난': 'angry',
    # READ-120-004 주제
    '공원에서 보낸 즐거운 하루': 'A fun day at the park',
    '아이스크림을 만드는 방법': 'How to make ice cream',
    '연을 만드는 방법': 'How to make a kite',
    '학교에 가는 길': 'The way to school',
    # READ-260-002 full
    '배가 부른 상태': 'having a full tummy',
    '배가 고픈 상태': 'feeling very hungry',
    '많이 아픈 상태': 'being very sick',
    '아주 기쁜 상태': 'feeling very happy',
    # READ-260 원인 (잭 양말)
    '장화를 신지 않아서': 'He did not wear his rain boots.',
    '노란 장화를 신어서': 'He wore his yellow rain boots.',
    '우산을 학교에 두고 와서': 'He left his umbrella at school.',
    '창문을 열어 두어서': 'He left the window open.',
    # READ-260 지호 편지
    '새 강아지 소식을 알리려고': 'To tell her about his new puppy',
    '새 소파를 자랑하려고': 'To show off his new sofa',
    '할머니의 생신을 축하하려고': 'To say happy birthday to Grandma',
    '밥 주는 방법을 물어보려고': 'To ask how to feed a dog',
    # READ-330 saving
    '(돈을) 모으고 있다': 'keeping money to use later',
    '(사람을) 구하고 있다': 'helping someone in danger',
    '(돈을) 쓰고 있다': 'spending money',
    '(돈을) 세고 있다': 'counting money',
    # READ-330 대회 수요일
    '선생님이 아프셔서': 'Because the teacher was sick',
    '댄이 우승을 해서': 'Because Dan won the contest',
    '학생들이 책을 읽지 않아서': 'Because the students did not read books',
    '월요일이 쉬는 날이어서': 'Because Monday was a holiday',
    # READ-330 벌
    '벌이 자연에서 하는 중요한 일': 'The important work bees do in nature',
    '꿀을 맛있게 먹는 방법': 'How to enjoy eating honey',
    '꽃을 예쁘게 기르는 방법': 'How to grow pretty flowers',
    '과일이 자라는 곳': 'Places where fruit grows',
    # READ-430 look after
    '돌보다': 'to take care of',
    '찾아보다': 'to look for',
    '쳐다보다': 'to look at',
    '따라가다': 'to follow',
    # READ-430 요지
    '작은 실천이 큰 변화를 만든다': 'Small actions can make a big change.',
    '쓰레기통을 깨끗이 비우는 방법': 'How to empty the trash can',
    '점심 도시락을 맛있게 싸는 방법': 'How to pack a tasty lunch',
    '운동장에서 재미있게 노는 방법': 'How to have fun on the playground',
    # READ-550 팬케이크
    '위에 작은 거품이 생겼을 때': 'When small bubbles appear on top',
    '팬을 2분 동안 데운 직후에': 'Right after heating the pan for two minutes',
    '꿀을 위에 올린 다음에': 'After putting honey on top',
    '팬케이크가 조각났을 때': 'When the pancake breaks into pieces',
    # READ-550 run out of
    '~을 다 써서 없다': 'to have no more of something',
    '~의 밖으로 뛰어나가다': 'to run outside quickly',
    '~을 바닥에 쏟다': 'to spill something on the floor',
    '~을 사러 나가다': 'to go out to buy something',
    # READ-550 소풍 장소
    '공원이 분수 공사로 문을 닫아서': 'The park closed for fountain repairs.',
    '화요일에 비가 많이 와서': 'It rained a lot on Tuesday.',
    '아침 날씨가 흐렸기 때문에': 'The morning sky was cloudy.',
    '학교 정원이 더 넓어서': 'The school garden is bigger.',
    # READ-550 다나 이메일
    '연습 횟수를 줄여도 되는지 부탁하려고': 'To ask if she can practice fewer times',
    '수영을 완전히 그만두겠다고 말하려고': 'To say she will quit swimming',
    '코치님께 감사 인사만 전하려고': 'To only say thank you to the coach',
    '바이올린 대회 소식을 알리려고': 'To share news about a violin contest',
    # READ-550 this habit
    '휴대폰을 계속 확인하는 것': 'Checking her phone all the time',
    '유리문에 부딪히는 것': 'Walking into the glass door',
    '계단에서 뛰어다니는 것': 'Running on the stairs',
    '아침에 일찍 일어나는 것': 'Waking up early in the morning',
    # READ-550 일치 (벤 트로피)
    '벤은 전에는 무언가를 이겨 본 적이 없다': 'Ben never won anything before.',
    '트로피는 벤 혼자서 노력한 결과이다': "The trophy was Ben's work alone.",
    '팀은 일 년 동안 벤을 기다려 주었다': 'The team waited for Ben for a year.',
    '벤은 트로피를 들고 눈물을 흘렸다': 'Ben cried while holding the trophy.',
    # READ-750 문어
    '세 번째 심장이 뛰는 것을 멈춘다': 'Its third heart stops beating.',
    '심장 두 개가 뛰는 것을 멈춘다': 'Two of its hearts stop beating.',
    '피가 파란색으로 변한다': 'Its blood turns blue.',
    '피부색을 바꿀 수 없게 된다': 'It cannot change its skin color.',
    # READ-750 홍수
    '강 밑에 쌓인 쓰레기가 물길을 막아서': 'Trash under the river blocked the water.',
    '오래된 다리가 너무 낮아서': 'The old bridge was too low.',
    '새 다리를 잘못 지어서': 'The new bridge was built wrong.',
    '여름마다 비가 너무 많이 와서': 'It rained too much every summer.',
    # READ-750 결론 (태블릿 vs 종이책)
    '상황에 맞게 둘을 골라 쓰는 법을 배워야 한다': 'Learn to pick the right one for each situation.',
    '태블릿이 종이책보다 언제나 낫다': 'Tablets are always better than paper books.',
    '눈 건강을 위해 종이책만 써야 한다': 'Use only paper books for healthy eyes.',
    '수업 중 게임은 공부에 도움이 된다': 'Games in class help students study.',
    # READ-750 편지 쓴 사람
    '할머니': 'Grandma',
    '엄마': 'Mom',
    '멀리서 공부하던 할머니': 'Grandma when she was studying far away',
    '글쓴이(나)': 'the writer (me)',
    # READ-750 순서 (유나)
    '유나가 연습 중에 발목을 다쳤다': 'Yuna hurt her ankle during practice.',
    '유나가 메달을 들고 무대에 섰다': 'Yuna stood on stage with a medal.',
    '동생이 유나를 링크장에 데려갔다': 'Her sister took her to the rink.',
    '관중이 유나의 이름을 외쳤다': "The crowd shouted Yuna's name.",
    # READ-750 의도 (장난감 편지)
    '빠진 부품을 보내 달라고 요청하려고': 'To ask the company to send a missing part',
    '장난감 회사의 디자인을 칭찬하려고': "To praise the toy company's design",
    '전체 금액을 환불해 달라고 요구하려고': 'To demand a full refund',
    '아들의 생일 파티에 초대하려고': "To invite them to his son's birthday party",
    # READ-900 공원 시원한 이유
    '나뭇잎이 수증기를 내보내며 공기를 식혀 주기 때문에': 'Leaves cool the air by releasing water vapor.',
    '키 큰 나무들의 그늘이 공기를 차갑게 만들기 때문에': 'Shade from tall trees makes the air cold.',
    '공원에는 바람이 늘 많이 불기 때문에': 'The park is always windy.',
    '사람들이 나무를 열심히 보호하기 때문에': 'People work hard to protect the trees.',
    # READ-900 제목
    '두 가지 문제가 만나 태어난 발명품': 'An Invention Born from Two Problems',
    '아이스크림을 시원하게 보관하는 방법': 'How to Keep Ice Cream Cold',
    '와플 장사가 어려웠던 이유': 'Why Selling Waffles Was Hard',
    '1904년 미국 박람회의 역사': 'The History of the 1904 Fair',
    # READ-900 일치 (도서관)
    '학기 마지막 날 전에 빌린 책을 모두 반납하면 다음 학기에 두 배로 빌릴 수 있다':
        'Return all borrowed books before the last day, and you can borrow twice as many next term.',
    '학기 마지막 날에 책을 두 권씩 반납해야 한다': 'Students must return two books each on the last day.',
    '다음 학기에는 책을 두 배 빠르게 읽어야 한다': 'Students must read twice as fast next term.',
    '교장 선생님이 도서관에 책을 두 배 많이 사 주었다': 'The principal bought twice as many books for the library.',
    # READ-900 콩 말리기
    '발효가 끝난 다음에': 'After the fermentation is finished',
    '나무에서 열매를 자르기 전에': 'Before cutting the pods from the tree',
    '공장에서 콩을 볶은 다음에': 'After the beans are roasted at the factory',
    '설탕과 섞는 것과 동시에': 'At the same time as mixing in the sugar',
    # READ-900 의도 (공원 회의)
    '시민들이 회의에 나와 공원 개선을 요구하게 하려고': 'To get people to come to the meeting and ask for a better park',
    '고속도로 꽃길의 아름다움을 소개하려고': 'To describe the beauty of the highway flowers',
    '시의 예산 사용을 객관적으로 설명하려고': 'To explain the city budget fairly',
    '도란초등학교의 위치를 알리려고': 'To show where Doran Elementary is',
}

assert len(SWAP) == 27 * 4, len(SWAP)

SRC = '/mnt/user-data/outputs/문제은행_리딩.xlsx'
wb = openpyxl.load_workbook(SRC)
ws = wb['리딩']
h = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(h)}
opt_cols = [col[f'보기{n}'] for n in range(1, 5)]

replaced = 0
touched_rows = set()
for row in ws.iter_rows(min_row=2):
    for ci in opt_cols:
        cell = row[ci - 1]
        v = str(cell.value or '').strip()
        if v in SWAP:
            cell.value = SWAP[v]
            replaced += 1
            touched_rows.add(row[0].row)

# 검증 1: 교체 수 = 27문항 × 4보기
assert replaced == 108, f'교체 수 이상: {replaced}'
assert len(touched_rows) == 27, f'문항 수 이상: {len(touched_rows)}'

# 검증 2: 보기 칸에 한국어가 하나도 남지 않음
ko = re.compile(r'[가-힣]')
for row in ws.iter_rows(min_row=2):
    for ci in opt_cols:
        v = str(row[ci - 1].value or '')
        assert not ko.search(v), f'{row[0].value} 보기에 한국어 잔존: {v}'

# 검증 3: 질문·지문·구멍 태그는 손대지 않음 (질문에 한국어 유지 확인)
q_ko = sum(1 for row in ws.iter_rows(min_row=2) if ko.search(str(row[col['질문'] - 1].value or '')))
assert q_ko == 44, q_ko
gap_cols = [col[f'보기{n}구멍'] for n in range(1, 5)]
gap_cnt = sum(1 for row in ws.iter_rows(min_row=2) for ci in gap_cols if str(row[ci - 1].value or '').strip())
assert gap_cnt == 44 * 3, gap_cnt  # 오답 3개마다 태그

wb.save(SRC)
print(f'보기 영어 전환 완료: {len(touched_rows)}문항 {replaced}칸 교체, 한국어 보기 0, 구멍 태그 {gap_cnt}개 유지')
