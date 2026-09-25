# -*- coding: utf-8 -*-
# 파닉스 판별형 문제은행 (2026-09-24 확정)
# - 단계당 4문항 × 5단계 = 20문항 (레테는 단계별 랜덤 2문항 출제 → 재응시 대비 여분 2개)
# - 각 단계 진짜 단어 2 + 무의미(외계어) 단어 2 — 아는 단어 찍기 방지, 소리 규칙만으로 풀게
# - 무의미 문항 질문에 "세상에 없는 외계어 단어" 안내 포함 (아이 당황 방지)
# - 전부 듣기형: 1단계 = 듣고 첫 글자, 2~5단계 = 듣고 맞는 철자
# - 음성파일 열 = TTS로 읽을 단어
# 단계: 1 글자 소리 / 2 단모음 CVC / 3 장모음(magic e·모음팀) / 4 자음 블렌드 / 5 이중글자(sh·ch·th)
import openpyxl
from openpyxl.styles import Font

Q_LETTER = '소리를 듣고 첫 글자를 고르세요.'
Q_WORD = '소리를 듣고 맞는 단어를 고르세요.'
ALIEN = ' (세상에 없는 외계어 단어예요!)'

# (단계, 질문, 정답, [오답3], TTS단어, 무의미여부)
Q = [
# ── 1단계: 첫 글자 소리 ──
(1, Q_LETTER, 'b', ['d', 'p', 'v'], 'ball', False),
(1, Q_LETTER, 's', ['z', 'c', 'j'], 'sun', False),
(1, Q_LETTER, 'v', ['f', 'w', 'b'], 'vop', True),   # v/f 무성·유성 혼동
(1, Q_LETTER, 'k', ['g', 'd', 't'], 'ked', True),   # k/g 혼동
# ── 2단계: 단모음 CVC ──
(2, Q_WORD, 'pen', ['pin', 'pan', 'pun'], 'pen', False),
(2, Q_WORD, 'hot', ['hat', 'hut', 'hit'], 'hot', False),
(2, Q_WORD, 'lut', ['lot', 'lit', 'let'], 'lut', True),
(2, Q_WORD, 'zam', ['zim', 'zom', 'zum'], 'zam', True),
# ── 3단계: 장모음·magic e ──
(3, Q_WORD, 'cake', ['cap', 'kick', 'coke'], 'cake', False),
(3, Q_WORD, 'rain', ['ran', 'run', 'rail'], 'rain', False),
(3, Q_WORD, 'mape', ['map', 'mep', 'mip'], 'mape', True),   # magic e 유무 판별
(3, Q_WORD, 'fode', ['fod', 'fade', 'food'], 'fode', True), # 장모음 소리 구분
# ── 4단계: 자음 블렌드 ──
(4, Q_WORD, 'stop', ['top', 'shop', 'sop'], 'stop', False),
(4, Q_WORD, 'frog', ['fog', 'log', 'flag'], 'frog', False),
(4, Q_WORD, 'snad', ['sad', 'sand', 'snap'], 'snad', True), # 블렌드 유무·자리
(4, Q_WORD, 'frip', ['fip', 'rip', 'flip'], 'frip', True),
# ── 5단계: 이중글자 sh·ch·th ──
(5, Q_WORD, 'ship', ['sip', 'chip', 'sheep'], 'ship', False),
(5, Q_WORD, 'think', ['sink', 'thin', 'pink'], 'think', False),
(5, Q_WORD, 'thop', ['top', 'shop', 'chop'], 'thop', True), # th/t/sh/ch 구분
(5, Q_WORD, 'shum', ['sum', 'chum', 'sham'], 'shum', True),
]

# ── 검증 ──
# 무의미 단어가 실제 영어 단어가 아닌지 스크리닝 (초등 노출 단어 사전 기준)
REAL_WORDS = set('''ball sun pen hot cake rain stop frog ship think
cat dog map cap mop tip top sad sand snap rip flip fog log flag
sip chip sheep sink thin pink sum chum sham fade food lot lit let
pin pan pun hat hut hit ran run rail kick coke shop chop com
fod mep mip zim zom zum fip'''.split())
NONSENSE = [t for (_, _, _, _, t, alien) in Q if alien]
for w in NONSENSE:
    assert w not in REAL_WORDS, f'무의미 단어가 진짜 단어와 겹침: {w}'
# 규칙 위반 없는 발음 가능 철자인지 (모음 포함, 3~5자)
for w in NONSENSE:
    assert any(v in w for v in 'aeiou') and 3 <= len(w) <= 5, w

levels, alien_cnt = {}, {}
for lv, q, ans, ds, tts, alien in Q:
    levels[lv] = levels.get(lv, 0) + 1
    if alien: alien_cnt[lv] = alien_cnt.get(lv, 0) + 1
    assert len(set([ans] + ds)) == 4, (q, ans)
    assert tts, ans
    if q == Q_WORD:
        assert ans == tts, ans           # 듣는 단어 = 정답 철자
    else:
        assert tts.startswith(ans), ans  # 첫 글자 문항: TTS 첫 글자 = 정답
assert len(levels) == 5 and all(v == 4 for v in levels.values()), levels
assert all(alien_cnt.get(s) == 2 for s in range(1, 6)), alien_cnt  # 단계마다 무의미 2개

# ── 엑셀 생성 (탭 통째 교체용 — 열 구조는 기존 문제은행_파닉스 탭과 동일) ──
wb = openpyxl.Workbook()
ws = wb.active
ws.title = '파닉스'
cols = ['문제ID', '게이트단계(1-5)', '질문유형', '질문', '보기1', '보기2', '보기3', '보기4', '정답번호', '이미지파일', '음성파일']
ws.append(cols)
for c in ws[1]:
    c.font = Font(name='Arial', bold=True)

counter, pos = {}, 0
for lv, q, ans, ds, tts, alien in Q:
    counter[lv] = counter.get(lv, 0) + 1
    qid = f'PHON-{lv}-{counter[lv]:03d}'
    question = q + (ALIEN if alien else '')
    answer_idx = pos % 4
    pos += 1
    opts = [None] * 4
    opts[answer_idx] = ans
    di = 0
    for i in range(4):
        if opts[i] is None:
            opts[i] = ds[di]
            di += 1
    ws.append([qid, lv, '듣고고르기', question] + opts + [answer_idx + 1, '', tts])

for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name='Arial')
ws.column_dimensions['D'].width = 52
ws.column_dimensions['A'].width = 14
ws.column_dimensions['C'].width = 12

wb.save('/mnt/user-data/outputs/문제은행_파닉스.xlsx')

from collections import Counter
apos = Counter(r[8] for r in ws.iter_rows(min_row=2, values_only=True))
print('문항 수:', ws.max_row - 1)
print('정답 위치 분포:', dict(sorted(apos.items())))
print('단계별 무의미:', alien_cnt)
print('전부 듣기형·음성 채움:', all(r[10] for r in ws.iter_rows(min_row=2, values_only=True)))
