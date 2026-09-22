# -*- coding: utf-8 -*-
"""지문 보강 패치 — 짧게 나온 문항에 내용 문장 추가 (문항 의도는 유지)"""
import openpyxl
ADD = {
 'READ-260-002': " He lies on the sofa and rubs his tummy.",
 'READ-260-006': " Mia felt very proud of her cookies.",
 'READ-330-002': " Her little brother thinks she is amazing.",
 'READ-330-003': " Leo promised to check his plant every single morning.",
 'READ-330-004': " Everyone clapped loudly for him.",
 'READ-330-006': " After the game, Ken washed the old cap very carefully.",
 'READ-430-002': " It was the best night of our summer.",
 'READ-430-004': (" Our captain lowered his head and said sorry to everyone.", 'Then our coach gathered us in a circle.'),
 'READ-550-002': " Dad laughed and said, \"Next time, let's buy two bags.\"",
 'READ-550-003': " His eyes looked young again while he spoke.",
 'READ-550-004': " Some of us even said the garden was better than the park.",
 'READ-550-006': " Her mom promised to put her own phone away at dinner, too.",
 'READ-550-007': " Reporters asked him how he felt, but he could only laugh.",
 'READ-750-001': " Some octopuses even carry coconut shells and use them as little houses.",
 'READ-750-002': " Now he bakes twice as many croissants every morning.",
 'READ-750-003': " The town still keeps the new bridge, but now it also cleans the river every spring.",
 'READ-750-004': " Some schools now use tablets for research time and paper books for reading time.",
 'READ-750-005': " That night, Mom called Grandma, and they talked for a very long time.",
 'READ-750-006': " She knew exactly who deserved to hold this medal with her.",
 'READ-750-007': " My son loved the picture of the silver robot on the front.",
 'READ-900-001': " The researchers now hope to plant more trees along the city's hottest bus stops.",
 'READ-900-002': " In one experiment, people who wrote directions by hand found their way faster than those who only followed their phones.",
 'READ-900-003': " Today, millions of ice cream cones are enjoyed around the world every single day.",
 'READ-900-004': " The rule counts every kind of book, from comic books to science books. Teachers say that even students who never visited the library are suddenly reading every day.",
 'READ-900-005': " Farmers taste a few beans from every pile to check that the fermentation went well.",
}
INSERT_BEFORE = {'READ-430-004', 'READ-750-007'}
ANCHOR = {'READ-430-004': 'Then our coach gathered us in a circle.',
          'READ-750-007': 'However, when we opened it,'}

wb = openpyxl.load_workbook('/mnt/user-data/outputs/레벨테스트_문제은행_완성본.xlsx')
ws = wb['리딩']
h = [c.value for c in ws[1]]
col = {name: i+1 for i, name in enumerate(h)}
patched = 0
for row in ws.iter_rows(min_row=2):
    qid = row[col['문제ID']-1].value
    if qid in ADD:
        cell = row[col['지문']-1]
        add = ADD[qid]
        if isinstance(add, tuple):
            add = add[0]
        if qid in INSERT_BEFORE:
            anchor = ANCHOR[qid]
            assert anchor in cell.value, qid
            cell.value = cell.value.replace(anchor, add.strip() + " " + anchor, 1)
        else:
            cell.value = cell.value.rstrip() + add
        patched += 1
wb.save('/mnt/user-data/outputs/레벨테스트_문제은행_완성본.xlsx')
print("보강", patched, "문항")
